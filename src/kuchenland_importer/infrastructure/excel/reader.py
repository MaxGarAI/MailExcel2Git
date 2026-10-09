"""Read OOXML values and formula-cache status without saving source workbooks."""

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import cast

from openpyxl import load_workbook

from kuchenland_importer.domain.cell_values import CellValue, ExcelErrorValue
from kuchenland_importer.domain.workbook import SourceCell, SourceSheet


def cell_value(value: object, is_error: bool = False) -> CellValue:
    if is_error:
        return ExcelErrorValue(str(value))
    if value is None or isinstance(value, str | int | float | bool | Decimal | date | datetime):
        return cast(CellValue, value)
    raise ValueError(f"Неподдерживаемый тип ячейки: {type(value).__name__}")


class WorkbookReader:
    def read(self, path: Path) -> tuple[SourceSheet, ...]:
        if path.suffix.lower() in {".xls", ".xlsb"}:
            from kuchenland_importer.infrastructure.excel.legacy_reader import read_legacy

            return read_legacy(path)
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("Неизвестный формат Excel.")
        if path.stat().st_size > 50 * 1024 * 1024:
            raise ValueError("Вложение превышает лимит чтения 50 МБ.")
        formulas = load_workbook(path, data_only=False, keep_links=False)
        try:
            values = load_workbook(path, data_only=True, keep_links=False)
            try:
                sheets = []
                for sheet in formulas:
                    if sheet.max_row > 100_000 or sheet.max_column > 512:
                        raise ValueError("Лист превышает лимит 100000 строк / 512 столбцов.")
                    cached = values[sheet.title]
                    rows = tuple(
                        tuple(
                            SourceCell(
                                cell_value(
                                    ""
                                    if cell.data_type == "f"
                                    and cached.cell(cell.row, cell.column).value is None
                                    and cached.cell(cell.row, cell.column).data_type in {"str", "s"}
                                    else cached.cell(cell.row, cell.column).value,
                                    cached.cell(cell.row, cell.column).data_type == "e",
                                ),
                                cell.number_format,
                                cell.data_type == "f"
                                and cached.cell(cell.row, cell.column).value is None
                                and cached.cell(cell.row, cell.column).data_type
                                not in {"str", "s"},
                            )
                            for cell in row
                        )
                        for row in sheet.iter_rows()
                    )
                    merges = tuple(
                        (r.min_row, r.min_col, r.max_row, r.max_col)
                        for r in sheet.merged_cells.ranges
                    )
                    sheets.append(SourceSheet(sheet.title, rows, merges))
                return tuple(sheets)
            finally:
                values.close()
        finally:
            formulas.close()
