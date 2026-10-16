"""Reopen the saved package and verify both mirrors, typed errors and native photo pixels."""

from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from kuchenland_importer.application.article_identity import article_value
from kuchenland_importer.application.target_fields import product_fields
from kuchenland_importer.domain.cell_values import CellValue, ExcelErrorValue
from kuchenland_importer.domain.master import MasterSnapshot
from kuchenland_importer.domain.workbook import SourceCell
from kuchenland_importer.domain.write_plan import WritePlan
from kuchenland_importer.infrastructure.excel.master_writer import WrittenProduct
from kuchenland_importer.infrastructure.excel.photo_reader import PhotoReader
from kuchenland_importer.infrastructure.excel.read_limits import MASTER_LIMITS
from kuchenland_importer.infrastructure.photo_store import PhotoStore


def equal_value(expected: CellValue, actual: object) -> bool:
    if expected is None or expected == "":
        return actual is None or actual == ""
    if isinstance(expected, Decimal):
        return isinstance(actual, float | int) and abs(float(expected) - actual) <= max(
            1e-9,
            abs(float(expected)) * 1e-12,
        )
    if isinstance(expected, datetime):
        return (
            isinstance(actual, datetime)
            and abs((expected.replace(tzinfo=None) - actual).total_seconds()) < 0.001
        )
    if isinstance(expected, date):
        return actual == expected or isinstance(actual, datetime) and actual.date() == expected
    if isinstance(expected, float):
        return isinstance(actual, float | int) and abs(expected - actual) <= max(
            1e-9,
            abs(expected) * 1e-12,
        )
    return expected == actual


def verify_master(
    path: Path,
    master: MasterSnapshot,
    plan: WritePlan,
    written: tuple[WrittenProduct, ...],
    directory: Path,
) -> None:
    verify_inventory(path, master, written)
    book = load_workbook(path, read_only=False, data_only=False)
    try:
        for action, saved in zip(plan.products, written, strict=True):
            for name, row, photo in (
                (saved.calculation_sheet, saved.calculation_row, False),
                (saved.photo_sheet, saved.photo_row, True),
            ):
                sheet = book[name]
                layout = master.layouts[name]
                for key, value in product_fields(action.incoming, photo).items():
                    if key not in layout.columns:
                        continue
                    cell = sheet.cell(row, layout.columns[key])
                    if isinstance(value, ExcelErrorValue):
                        valid = cell.data_type == "e" and cell.value == value.code
                    else:
                        valid = cell.data_type != "f" and equal_value(value, cell.value)
                    if not valid:
                        raise ValueError(
                            f"Проверка записи не пройдена: {name}!{cell.coordinate}; "
                            f"ожидалось {value!r}, получено {cell.value!r}"
                        )
    finally:
        book.close()
    extraction = PhotoReader().read(path, limits=MASTER_LIMITS)
    photos = {(p.sheet, p.row, p.column): p for p in extraction.pictures if p.native}
    store = PhotoStore(directory)
    linked = {
        (p.sheet, p.row, p.column, p.shape_name)
        for p in extraction.problems if p.external_link
    }
    for saved in written:
        column = master.layouts[saved.photo_sheet].columns["photo"]
        if any(
            (saved.photo_sheet, saved.photo_row, column, name) not in linked
            for name in saved.linked_photo_names
        ):
            raise ValueError(f"Прежняя внешняя ссылка на фото не сохранена: {saved.article}")
        if saved.photo_digest is None:
            continue
        column = master.layouts[saved.photo_sheet].columns["photo"]
        picture = photos.get((saved.photo_sheet, saved.photo_row, column))
        if picture is None or store.build((picture,)).pixel_sha256 != saved.photo_digest:
            raise ValueError(f"Нативное фото не прошло проверку: {saved.article}")


def verify_inventory(
    path: Path, master: MasterSnapshot, written: tuple[WrittenProduct, ...]
) -> None:
    """Verify every old article/mirror remains, except explicitly replaced product rows."""
    replaced = {product.article for product in written}
    expected = Counter(
        (article, row.sheet)
        for article, product in master.products.items() if article not in replaced
        for row in product.rows
    )
    for saved in written:
        expected.update(
            ((saved.article, saved.calculation_sheet), (saved.article, saved.photo_sheet))
        )
    actual: Counter[tuple[str, str]] = Counter()
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for name, layout in master.layouts.items():
            column = layout.columns["article"]
            for (cell,) in book[name].iter_rows(
                min_row=layout.first_row, min_col=column, max_col=column
            ):
                if cell.value is not None and cell.value != "":
                    actual[(article_value(SourceCell(cell.value)), name)] += 1
    finally:
        book.close()
    if actual != expected:
        missing = sum((expected - actual).values())
        extra = sum((actual - expected).values())
        raise ValueError(
            f"Проверка сохранности артикулов не пройдена: отсутствует строк {missing}, "
            f"лишних строк {extra}. Рабочая копия не будет принята."
        )
