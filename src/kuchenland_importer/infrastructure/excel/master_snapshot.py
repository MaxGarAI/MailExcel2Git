"""Read a closed master, retain excluded tabs, and resolve existing photos only when used."""

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from kuchenland_importer.application.normalize_workbook import article_value
from kuchenland_importer.application.table_schema import ColumnRegistry, find_header
from kuchenland_importer.domain.master import ExistingProduct, ExistingRow, MasterSnapshot
from kuchenland_importer.domain.photos import EmbeddedPhoto
from kuchenland_importer.domain.season import SeasonRoute
from kuchenland_importer.infrastructure.excel.master_layout import layout_for, master_registry
from kuchenland_importer.infrastructure.excel.photo_reader import PhotoReader
from kuchenland_importer.infrastructure.excel.reader import WorkbookReader
from kuchenland_importer.infrastructure.photo_store import PhotoStore

STATE_SHEET = "_KL_IMPORT_STATE"


def read_snapshot(
    path: Path,
    registry: ColumnRegistry,
    routes: tuple[SeasonRoute, ...],
    excluded: tuple[str, ...],
    photo_directory: Path,
) -> MasterSnapshot:
    sheets = WorkbookReader().read(path)
    configured = {name for route in routes for name in (route.calculation_sheet, route.photo_sheet)}
    excluded_keys = {name.casefold() for name in excluded}
    layouts = {}
    products: dict[str, ExistingProduct] = {}
    cells = {}
    for sheet in sheets:
        if sheet.name == STATE_SHEET or sheet.name.casefold() in excluded_keys:
            continue
        if sheet.name not in configured and find_header(sheet, master_registry(registry)) is None:
            continue
        layout = layout_for(sheet, registry)
        layout.calculation_sheet = next(
            (route.calculation_sheet for route in routes if route.photo_sheet == sheet.name),
            None,
        )
        layouts[sheet.name] = layout
        for row in range(layout.first_row, len(sheet.rows) + 1):
            cell = sheet.rows[row - 1][layout.columns["article"] - 1]
            if cell.value is None or cell.value == "":
                continue
            article = article_value(cell)
            item = products.setdefault(article, ExistingProduct([]))
            item.rows.append(ExistingRow(sheet.name, row))
            if "mail_received" in layout.columns:
                received = sheet.rows[row - 1][layout.columns["mail_received"] - 1].value
                day = received.date() if isinstance(received, datetime) else received
                if isinstance(day, date):
                    item.received_day = max(item.received_day or day, day)
            if "photo" in layout.columns:
                cells[(sheet.name, row, layout.columns["photo"])] = article
    extracted = PhotoReader().read(path, photo_directory)
    grouped: dict[str, list[EmbeddedPhoto]] = defaultdict(list)
    for photo in extracted.pictures:
        photo_article = cells.get((photo.sheet, photo.row, photo.column))
        if photo_article is None:
            continue
        if photo.end_row is not None and any(
            cells.get((photo.sheet, row, photo.column)) not in {None, photo_article}
            for row in range(photo.row + 1, photo.end_row + 1)
        ):
            products[photo_article].photo_error = "Прежнее фото пересекает строки разных товаров."
        else:
            grouped[photo_article].append(photo)
    for problem in extracted.problems:
        for (sheet_name, row, column), article in cells.items():
            if (
                sheet_name == problem.sheet
                and problem.row in {None, row}
                and (problem.column in {None, column})
            ):
                products[article].photo_error = problem.message
    store = PhotoStore(photo_directory)
    for article, pictures in grouped.items():
        try:
            products[article].photo = store.build(tuple(pictures))
        except Exception as error:
            products[article].photo_error = str(error)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        mode = book.calculation.calcMode if book.calculation is not None else None
        if STATE_SHEET in book.sheetnames:
            state = book[STATE_SHEET]
            if state.cell(1, 1).value != "kuchenland-import-state-v1":
                raise ValueError("Служебный лист состояния имеет неизвестный формат.")
            for article, received, fingerprint, *_ in state.iter_rows(min_row=2, values_only=True):
                if article in products:
                    parsed = datetime.fromisoformat(str(received))
                    if parsed.tzinfo is None:
                        raise ValueError("Дата служебного состояния без часового пояса.")
                    products[str(article)].received_at = parsed
                    products[str(article)].fingerprint = str(fingerprint)
    finally:
        book.close()
    return MasterSnapshot(layouts, products, tuple(sheet.name for sheet in sheets), mode or "auto")
