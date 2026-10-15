"""Regression: tiny attachment into a real photo workbook above the old 50 MiB cap."""

import argparse
import hashlib
import io
import json
import os
import random
import shutil
import subprocess
import sys
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as DrawingImage  # noqa: E402
from PIL import Image  # noqa: E402

from kuchenland_importer.domain.capture import MailCapture  # noqa: E402
from kuchenland_importer.domain.mail import MailMetadata, SavedAttachment  # noqa: E402
from kuchenland_importer.domain.season import SeasonRoute  # noqa: E402
from kuchenland_importer.infrastructure.capture_manifest import (
    JsonCaptureManifestStore,  # noqa: E402
)
from kuchenland_importer.infrastructure.config_loader import load_settings  # noqa: E402
from kuchenland_importer.infrastructure.settings_store import save_settings  # noqa: E402


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    parser.add_argument("--exe", type=Path, required=True)
    args = parser.parse_args()
    work, exe = args.directory.resolve(), args.exe.resolve()
    work.mkdir(parents=True, exist_ok=False)
    root = Path(__file__).resolve().parents[1]
    book = Workbook()
    calc = book.active
    calc.title = "Весна"
    photos = book.create_sheet("ФОТО весна")
    calc.append(["Артикул", "Поставщик", "Сезон", "Наименование"])
    photos.append(["Артикул", "Поставщик", "Сезон", "Наименование", "Фото"])
    for index in range(17):
        values = [f"OLD-{index}", "Vendor", "Весна", "Existing"]
        calc.append(values)
        photos.append(values)
        photos.row_dimensions[index + 2].height = 100
        stream = io.BytesIO()
        pixels = random.Random(index).randbytes(1024 * 1024 * 3)
        Image.frombytes("RGB", (1024, 1024), pixels).save(stream, format="PNG")
        picture = DrawingImage(stream)
        picture.width = picture.height = 80
        photos.add_image(picture, f"E{index + 2}")
    target = work / "master.xlsx"
    book.save(target)
    book.close()
    initial_size = target.stat().st_size
    assert initial_size > 50 * 1024 * 1024
    incoming = work / "source"
    incoming.mkdir()
    source = incoming / "small.xlsx"
    book = Workbook()
    book.active.append(["Артикул", "Поставщик", "Сезон", "Наименование", "Фото"])
    book.active.append(["NEW", "Vendor", "Весна", "Incoming"])
    stream = io.BytesIO()
    Image.new("RGB", (32, 32), "green").save(stream, format="PNG")
    book.active.add_image(DrawingImage(stream), "E2")
    book.save(source)
    book.close()
    source_hash = digest(source)
    mail = MailMetadata("qa", "qa", "Large-master regression", datetime.now(UTC), "QA", "")
    attachment = SavedAttachment(mail, 1, source.name, source, source_hash)
    manifest = JsonCaptureManifestStore().save(
        MailCapture("large-master", incoming, 1, (mail,), (attachment,), (), (1,))
    )
    config = work / "config"
    config.mkdir()
    for name in ("columns.toml", "writing.toml"):
        shutil.copy2(root / "config" / name, config / name)
    settings = replace(
        load_settings(root / "config/default.toml"),
        workbook=target,
        data_dir=work / "data",
        routes=(SeasonRoute("spring", ("Весна",), "Весна", "ФОТО весна"),),
    )
    save_settings(config / "local.toml", settings)
    environment = dict(
        os.environ,
        LOCALAPPDATA=str(work / "profile"),
        PATH=str(Path(os.environ["SystemRoot"]) / "System32"),
    )
    for name in ("PYTHONPATH", "PYTHONHOME", "VIRTUAL_ENV"):
        environment.pop(name, None)
    command = [
        str(exe),
        "import-excel",
        "--config",
        str(config / "local.toml"),
        "--manifest",
        str(manifest),
    ]
    with (work / "import-output.txt").open("wb") as output:
        result = subprocess.run(
            command, cwd=work, env=environment, stdout=output, stderr=subprocess.STDOUT, check=False
        )
    assert result.returncode == 0, f"Inspect {work / 'import-output.txt'}"
    report = next((work / "data/reports").glob("import-*.json"))
    data = json.loads(report.read_text(encoding="utf-8"))
    assert data["committed"] and data["report"]["products_added"] == 1
    assert data["report"]["photos_added"] == 1
    assert target.stat().st_size > 50 * 1024 * 1024
    assert digest(source) == source_hash
    saved = digest(target)
    with (work / "repeat-output.txt").open("wb") as output:
        result = subprocess.run(
            command, cwd=work, env=environment, stdout=output, stderr=subprocess.STDOUT, check=False
        )
    assert result.returncode == 0 and digest(target) == saved
    print(
        json.dumps(
            {
                "master_before_bytes": initial_size,
                "master_after_bytes": target.stat().st_size,
                "attachment_bytes": source.stat().st_size,
                "added": 1,
                "photos_added": 1,
                "repeat_unchanged": True,
            }
        )
    )


if __name__ == "__main__":
    main()
