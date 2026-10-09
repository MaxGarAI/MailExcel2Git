"""Disposable Excel acceptance: move, native photos, errors, filters, repeat and recovery."""

import argparse
import sys
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from shutil import copy2

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.drawing.image import Image as DrawingImage  # noqa: E402
from openpyxl.utils.datetime import CALENDAR_MAC_1904  # noqa: E402
from PIL import Image  # noqa: E402

from kuchenland_importer.application.plan_import import make_plan  # noqa: E402
from kuchenland_importer.application.resolve_season import ResolveSeason  # noqa: E402
from kuchenland_importer.domain.cell_values import ExcelErrorValue  # noqa: E402
from kuchenland_importer.domain.import_batch import PreparedAttachment  # noqa: E402
from kuchenland_importer.domain.mail import MailMetadata, SavedAttachment  # noqa: E402
from kuchenland_importer.domain.photos import (
    EmbeddedPhoto,  # noqa: E402
    PhotoAction,  # noqa: E402
)
from kuchenland_importer.domain.product import ProductRecord  # noqa: E402
from kuchenland_importer.domain.season import RoutedProduct, SeasonRoute  # noqa: E402
from kuchenland_importer.infrastructure.backup_recovery import restore_backup  # noqa: E402
from kuchenland_importer.infrastructure.excel.columns import load_columns  # noqa: E402
from kuchenland_importer.infrastructure.excel.com_values import (  # noqa: E402
    ERROR_FORMULAS,
    write_value,
)
from kuchenland_importer.infrastructure.excel.master_snapshot import read_snapshot  # noqa: E402
from kuchenland_importer.infrastructure.excel.master_writer import write_master  # noqa: E402
from kuchenland_importer.infrastructure.excel.owned_session import owned_workbook  # noqa: E402
from kuchenland_importer.infrastructure.excel.photo_reader import PhotoReader  # noqa: E402
from kuchenland_importer.infrastructure.excel.verify_master import verify_master  # noqa: E402
from kuchenland_importer.infrastructure.file_transaction import (  # noqa: E402
    WorkbookTransaction,
    file_hash,
)
from kuchenland_importer.infrastructure.photo_store import PhotoStore  # noqa: E402


def fixture(root):
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904
    book.remove(book.active)
    red = root / "red.png"
    Image.new("RGB", (50, 30), "red").save(red)
    for name, photo in (
        ("Весна", False),
        ("ФОТО весна", True),
        ("Осень", False),
        ("ФОТО осень", True),
    ):
        sheet = book.create_sheet(name)
        sheet["A1"] = "Служебная строка"
        sheet["E2"] = "=SUM(E4:E20)"
        headers = [
            "Артикул",
            "Фото" if photo else "Поставщик",
            "Сезон",
            "Наименование",
            "Цена",
            "Старт продаж",
            "Ручная заметка",
            "Служебная формула",
            "дата",
        ]
        for column, value in enumerate(headers, 1):
            sheet.cell(3, column, value)
        sheet.auto_filter.ref = "A3:I5"
        if "весна" in name.casefold():
            for row, sku in ((4, "A"), (5, "KEEP")):
                sheet.cell(row, 1, sku)
                if not photo:
                    sheet.cell(row, 2, "Vendor")
                sheet.cell(row, 3, "Весна")
                sheet.cell(row, 4, sku + " name")
                sheet.cell(row, 5, 10)
                sheet.cell(row, 7, "photo-keep" if photo else "calc-keep")
                sheet.cell(row, 8, f"=VLOOKUP(A{row},Весна!A:D,4,FALSE)" if photo else f"=E{row}*2")
                sheet.row_dimensions[row].height = 80
            if photo:
                picture = DrawingImage(red)
                sheet.add_image(picture, "B4")
        sheet.column_dimensions["B"].width = 25
    deleted = book.create_sheet("удалено")
    deleted["A1"] = "Не изменять"
    deleted["B2"] = "=1+2"
    target = root / "master.xlsx"
    book.save(target)
    book.close()
    return target


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    root = parser.parse_args().directory.resolve()
    root.mkdir(parents=True, exist_ok=False)
    errors_file = root / "error-values.xlsx"
    empty = Workbook()
    empty.save(errors_file)
    empty.close()
    with owned_workbook(errors_file) as excel:
        for row, code in enumerate(ERROR_FORMULAS, 1):
            write_value(excel.Worksheets(1).Cells(row, 1), ExcelErrorValue(code))
        excel.Save()
    errors_book = load_workbook(errors_file, data_only=False)
    assert [(c.value, c.data_type) for row in errors_book.active for c in row] == [
        (code, "e") for code in ERROR_FORMULAS
    ]
    errors_book.close()
    target = fixture(root)
    original = file_hash(target)
    routes = (
        SeasonRoute("spring", ("Весна",), "Весна", "ФОТО весна"),
        SeasonRoute("autumn", ("Осень",), "Осень", "ФОТО осень"),
    )
    registry = load_columns(Path(__file__).resolve().parents[1] / "config/columns.toml")
    store = PhotoStore(root / "photos")
    green = root / "green.png"
    Image.new("RGB", (50, 30), "green").save(green)
    asset = store.build((EmbeddedPhoto("Фото", 4, 2, green.read_bytes(), "QA"),))
    batch = []
    for sku, season, image in (("A", "Осень", None), ("B", "Весна", asset)):
        mail = MailMetadata(
            sku, "QA", "=literal subject", datetime(2026, 10, 9, 12, tzinfo=UTC), "", ""
        )
        attachment = SavedAttachment(mail, 1, "qa.xlsx", target, original)
        product = ProductRecord(
            sku,
            "Vendor",
            season,
            attachment,
            "Товары",
            2,
            {
                "article": sku,
                "supplier": "Vendor",
                "season": season,
                "name": "=literal name",
                "purchase_price": ExcelErrorValue("#DIV/0!"),
                "sales_start": date(2027, 2, 2),
                "source:extra": "skip",
            },
            image,
        )
        routed = RoutedProduct(product, ResolveSeason(routes).execute(product))
        batch.append(PreparedAttachment(attachment, (routed,), frozenset({sku}), ()))
    with WorkbookTransaction(target, root / "backup") as transaction:
        master = read_snapshot(transaction.work, registry, routes, ("удалено",), root / "old")
        plan = make_plan(tuple(batch), master)
        assert len(plan.products) == 2
        written = write_master(transaction.work, master, plan, preserve_manual=True)
        verify_master(transaction.work, master, plan, written, root / "verified")
        transaction.commit()
    master = read_snapshot(target, registry, routes, ("удалено",), root / "reopened")
    assert make_plan(tuple(batch), master).unchanged == 2
    assert {r.sheet for r in master.products["A"].rows} == {"Осень", "ФОТО осень"}
    book = load_workbook(target, data_only=False)
    assert book["ФОТО осень"]["G4"].value == "photo-keep"
    assert "Осень!" in book["ФОТО осень"]["H4"].value or "'Осень'!" in (
        book["ФОТО осень"]["H4"].value
    )
    assert book["Осень"]["G4"].value == "calc-keep"
    assert book["ФОТО весна"].auto_filter.ref == "A3:I5"
    assert book["Весна"]["E2"].value == "=SUM(E4:E20)"
    assert book["удалено"]["B2"].value == "=1+2"
    book.close()
    sorted_copy = root / "sorted.xlsx"
    copy2(target, sorted_copy)
    with owned_workbook(sorted_copy) as excel:
        sheet = excel.Worksheets("ФОТО весна")
        sorter = sheet.Sort
        sorter.SortFields.Clear()
        sorter.SortFields.Add(Key=sheet.Range("A4:A5"), SortOn=0, Order=1, DataOption=0)
        sorter.SetRange(sheet.Range("A3:I5"))
        sorter.Header = 1
        sorter.Orientation = 1  # xlTopToBottom for the Sort object.
        sorter.Apply()
        sheet.Range("A3:I5").AutoFilter(Field=1, Criteria1="B")
        excel.Save()
    sorted_book = load_workbook(sorted_copy, data_only=True)
    native = PhotoReader().read(sorted_copy).pictures
    found = [
        p
        for p in native
        if p.native and p.sheet == "ФОТО весна" and sorted_book[p.sheet].cell(p.row, 1).value == "B"
    ]
    assert len(found) == 1 and store.build(tuple(found)).pixel_sha256 == asset.pixel_sha256
    assert found[0].row == 4
    assert not sorted_book["ФОТО весна"].row_dimensions[4].hidden
    assert sorted_book["ФОТО весна"].row_dimensions[5].hidden
    sorted_book.close()
    blue = root / "blue.png"
    Image.new("RGB", (50, 30), "blue").save(blue)
    replacement = store.build((EmbeddedPhoto("Фото", 4, 2, blue.read_bytes(), "QA"),))
    newer_batch = []
    for item in batch:
        product = item.products[0].product
        mail = replace(
            product.attachment.mail,
            received_at=product.attachment.mail.received_at + timedelta(days=1),
        )
        product = replace(
            product,
            attachment=replace(product.attachment, mail=mail),
            photo=replacement if product.article == "A" else asset,
        )
        routed = RoutedProduct(product, ResolveSeason(routes).execute(product))
        newer_batch.append(PreparedAttachment(product.attachment, (routed,), item.articles, ()))
    with WorkbookTransaction(target, root / "backup") as update:
        next_plan = make_plan(tuple(newer_batch), master)
        changed = write_master(update.work, master, next_plan, preserve_manual=True)
        verify_master(update.work, master, next_plan, changed, root / "replaced")
        assert {p.article: p.photo_action for p in changed} == {
            "A": PhotoAction.REPLACE,
            "B": PhotoAction.KEEP,
        }
        update.commit()
    restore_backup(update.journal, target, root / "backup")
    previous = restore_backup(transaction.journal, target, root / "backup")
    assert previous is not None and file_hash(target) == original
    print(
        "PASS: move, absent-photo preservation, native sort/filter, errors, dates, repeat, backup."
    )


if __name__ == "__main__":
    main()
