import io
import json
from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as DrawingImage
from PIL import Image

from kuchenland_importer.application.resolve_season import ResolveSeason
from kuchenland_importer.application.select_latest import SelectLatest
from kuchenland_importer.domain.season import is_excluded_season
from kuchenland_importer.infrastructure.config_loader import load_settings
from kuchenland_importer.infrastructure.excel.columns import load_columns
from kuchenland_importer.infrastructure.prepare_batch import prepare_batch
from kuchenland_importer.infrastructure.preview_report import create_preview

ROOT = Path(__file__).resolve().parents[2]


@pytest.mark.parametrize(
    "season",
    [
        "ВСЕСЕЗОННЫЙ",
        " всесезонный ",
        "(Д) 9. ВСЕСЕЗОННЫЙ 2027",
        "ВСЕСЕЗОННЫЙ\u00a02027",
    ],
)
def test_excluded_season_accepts_observed_case_service_prefix_and_year(season):
    assert is_excluded_season(season)


@pytest.mark.parametrize(
    "season", [None, "Весна", "Всесезонный набор", "Не всесезонный", "НГ 2027"]
)
def test_exclusion_does_not_match_other_or_ambiguous_seasons(season):
    assert not is_excluded_season(season)


def mixed_capture(tmp_path):
    source = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.append(["Article", "Supplier", "Season", "Price"])
    book.active.append(["G", "Vendor", "Весна 2027", 5])
    # These ignored row errors must not make the other product ineligible.
    book.active.append(["X", None, "ВСЕСЕЗОННЫЙ", "=1+1"])
    photo = book.create_sheet("Images")
    photo.append(["Article", "Photo", "Name"])
    photo.append(["G", None, "Good"])
    photo.append(["X", None, "=1+1"])
    for row in (2, 3):
        photo.row_dimensions[row].height = 100
        buffer = io.BytesIO()
        Image.new("RGB", (32, 32), "green").save(buffer, format="PNG")
        photo.add_image(DrawingImage(buffer), f"B{row}")
    book.save(source)
    book.close()
    with ZipFile(source) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    parts["xl/media/image2.png"] = b"broken ignored image"
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "stage": "mail_capture",
                "mails": [
                    {
                        "entry_id": "entry",
                        "store_id": "store",
                        "subject": "Subject",
                        "received_at": datetime(2026, 10, 20, tzinfo=UTC).isoformat(),
                        "sender_name": "",
                        "sender_address": "",
                    }
                ],
                "attachments": [
                    {
                        "mail_index": 0,
                        "attachment_index": 1,
                        "original_name": source.name,
                        "path": source.name,
                        "sha256": sha256(source.read_bytes()).hexdigest(),
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    return source, manifest


def test_ignored_row_bad_data_and_photo_do_not_block_regular_product(tmp_path):
    source, manifest = mixed_capture(tmp_path)
    before = source.read_bytes()
    registry = load_columns(ROOT / "config/columns.toml")
    resolver = ResolveSeason(load_settings(ROOT / "config/default.toml").routes)
    batch = prepare_batch(manifest, registry, resolver, tmp_path / "photos")
    assert len(batch) == 1 and batch[0].eligible
    assert [p.product.article for p in batch[0].products] == ["G"]
    assert batch[0].products[0].product.photo is not None
    assert batch[0].ignored[0].article == "X"
    selected = SelectLatest().execute(batch)
    assert [p.product.article for p in selected.products] == ["G"]
    assert selected.skipped_articles == 1
    assert not any(i.severity == "error" for i in selected.issues)
    report = tmp_path / "preview.json"
    result = create_preview(
        manifest,
        ROOT / "config/columns.toml",
        report,
        season_resolver=resolver,
        photo_directory=tmp_path / "preview-photos",
    )
    data = json.loads(report.read_text(encoding="utf-8"))
    assert result.eligible_products == 1 and result.errors == 0 and result.photos == 1
    assert data["ignored_product_count"] == 1
    assert data["files"][0]["ignored_products"][0]["article"] == "X"
    assert source.read_bytes() == before


@pytest.mark.parametrize("article", ["X", None])
def test_all_excluded_rows_need_neither_valid_supplier_nor_configured_target_sheets(
    tmp_path, article
):
    from kuchenland_importer.domain.season import SeasonRoute

    source = tmp_path / "ignored.xlsx"
    book = Workbook()
    book.active.append(["Article", "Supplier", "Season", "Price"])
    book.active.append([article, None, "ВСЕСЕЗОННЫЙ", "=1+1"])
    book.save(source)
    book.close()
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "stage": "mail_capture",
                "mails": [
                    {
                        "entry_id": "entry",
                        "store_id": "store",
                        "subject": "Subject",
                        "received_at": datetime(2026, 10, 20, tzinfo=UTC).isoformat(),
                        "sender_name": "",
                        "sender_address": "",
                    }
                ],
                "attachments": [
                    {
                        "mail_index": 0,
                        "attachment_index": 1,
                        "original_name": source.name,
                        "path": source.name,
                        "sha256": sha256(source.read_bytes()).hexdigest(),
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    # Even an alias configured by the user cannot enable the permanently excluded category.
    resolver = ResolveSeason((SeasonRoute("all_year", ("ВСЕСЕЗОННЫЙ",), "All", "Photos"),))
    batch = prepare_batch(
        manifest, load_columns(ROOT / "config/columns.toml"), resolver, tmp_path / "photos"
    )
    assert batch[0].eligible and not batch[0].products
    selected = SelectLatest().execute(batch)
    assert not selected.products and selected.skipped_articles == 1
    assert not any(i.severity == "error" for i in selected.issues)
