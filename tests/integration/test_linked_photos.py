import io
from datetime import UTC, datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as DrawingImage
from PIL import Image

from kuchenland_importer.application.bind_photos import BindPhotos
from kuchenland_importer.application.normalize_workbook import NormalizeWorkbook
from kuchenland_importer.domain.mail import MailMetadata, SavedAttachment
from kuchenland_importer.domain.season import SeasonRoute
from kuchenland_importer.infrastructure.excel.columns import load_columns
from kuchenland_importer.infrastructure.excel.master_snapshot import read_snapshot
from kuchenland_importer.infrastructure.excel.ooxml_package import A, R
from kuchenland_importer.infrastructure.excel.photo_reader import PhotoReader
from kuchenland_importer.infrastructure.excel.reader import WorkbookReader
from kuchenland_importer.infrastructure.photo_store import PhotoStore

CONFIG = Path(__file__).resolve().parents[2] / "config/columns.toml"


def linked_book(path: Path, *, embedded=False, broken=False):
    book = Workbook()
    book.active.title = "Items"
    book.active.append(["Article", "Supplier", "Season"])
    book.active.append(["001", "Vendor", "Весна 2027"])
    sheet = book.create_sheet("Images")
    sheet.append(["Article", "Photo", "Name"])
    sheet.append(["001", None, "First"])
    sheet.row_dimensions[2].height = 100
    buffer = io.BytesIO()
    Image.new("RGB", (40, 40), "blue").save(buffer, format="PNG")
    sheet.add_image(DrawingImage(buffer), "B2")
    book.save(path)
    book.close()
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    drawing = ET.fromstring(parts["xl/drawings/drawing1.xml"])
    blip = drawing.find(f".//{{{A}}}blip")
    if not embedded:
        blip.attrib.pop(f"{{{R}}}embed")
    blip.set(f"{{{R}}}link", "missing" if broken else "external")
    parts["xl/drawings/drawing1.xml"] = ET.tostring(drawing)
    rels = ET.fromstring(parts["xl/drawings/_rels/drawing1.xml.rels"])
    ET.SubElement(
        rels,
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
        Id="external",
        Type=R + "/image",
        Target="https://example.invalid/image.jpg",
        TargetMode="External",
    )
    parts["xl/drawings/_rels/drawing1.xml.rels"] = ET.tostring(rels)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


@pytest.mark.parametrize("embedded,broken", [(False, False), (True, False), (False, True)])
def test_link_without_bytes_is_warning_but_broken_relationship_is_error(tmp_path, embedded, broken):
    path = tmp_path / "source.xlsx"
    linked_book(path, embedded=embedded, broken=broken)
    before = path.read_bytes()
    extracted = PhotoReader().read(path)
    registry = load_columns(CONFIG)
    sheets = WorkbookReader().read(path)
    mail = MailMetadata("entry", "store", "Subject", datetime(2026, 10, 16, tzinfo=UTC), "", "")
    attachment = SavedAttachment(mail, 1, path.name, path, "0" * 64)
    preview = NormalizeWorkbook(registry).execute(attachment, sheets)
    result = BindPhotos(registry, PhotoStore(tmp_path / "photos")).execute(
        preview, sheets, extracted
    )
    assert len(result.products) == 1
    errors = [i for i in result.issues if i.severity == "error"]
    assert bool(errors) == broken
    assert bool(result.products[0].photo) == embedded
    if not embedded and not broken:
        assert any(
            i.code == "PHOTO_EXTERNAL_LINK" and i.severity == "warning" for i in result.issues
        )
    assert path.read_bytes() == before


def test_master_retains_link_identity_without_calling_it_corrupt(tmp_path):
    path = tmp_path / "master.xlsx"
    linked_book(path)
    snapshot = read_snapshot(
        path,
        load_columns(CONFIG),
        (SeasonRoute("spring", ("Весна",), "Items", "Images"),),
        (),
        tmp_path / "photos",
    )
    product = snapshot.products["001"]
    assert product.photo is None and product.photo_error is None
    assert len(product.linked_photos) == 1
    assert product.linked_photos[0].sheet == "Images"


@pytest.mark.parametrize("existing_drawing", [False, True])
def test_external_picture_survives_package_roundtrip_after_excel_rewrites_sheet(
    tmp_path, existing_drawing
):
    from types import SimpleNamespace

    from kuchenland_importer.domain.master import ExistingProduct, LinkedPhotoReference
    from kuchenland_importer.infrastructure.excel.linked_photo_keeper import LinkedPhotoKeeper

    path = tmp_path / "master.xlsx"
    linked_book(path)
    extracted = PhotoReader().read(path)
    reference = LinkedPhotoReference("Images", extracted.problems[0].shape_name)
    old = ExistingProduct([], linked_photos=(reference,))
    action = SimpleNamespace(
        previous=old, incoming=SimpleNamespace(product=SimpleNamespace(article="001", photo=None))
    )
    keeper = LinkedPhotoKeeper(path, SimpleNamespace(products=(action,)))
    assert not PhotoReader().read(path).problems  # Excel receives no external image anchor.
    if existing_drawing:
        linked_book(path, embedded=True)
    else:
        book = Workbook()
        book.active.title = "Destination"
        book.active.append(["Article", "Photo", "Name"])
        book.active.append(["001", None, "Updated"])
        book.save(path)
        book.close()
    cell = SimpleNamespace(RowHeight=100, Width=100, Height=100, Left=100, Top=20, Column=2, Row=2)
    names = keeper.restore(
        "001", SimpleNamespace(Name="Images" if existing_drawing else "Destination"), cell
    )
    keeper.finish(path)
    result = PhotoReader().read(path)
    external = [p for p in result.problems if p.external_link]
    assert len(external) == 1
    assert (external[0].row, external[0].column) == (2, 2)
    assert external[0].shape_name == names[0]
    assert len(result.pictures) == int(existing_drawing)
    with ZipFile(path) as archive:
        relationships = b"".join(archive.read(n) for n in archive.namelist() if n.endswith(".rels"))
        assert b"https://example.invalid/image.jpg" in relationships
