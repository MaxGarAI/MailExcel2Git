import io
import json
from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from kuchenland_importer.infrastructure.excel.ooxml_package import D
from kuchenland_importer.infrastructure.excel.photo_reader import PhotoReader
from kuchenland_importer.presentation.cli import main

CONFIG = Path(__file__).resolve().parents[2] / "config" / "default.toml"


@pytest.mark.parametrize("collapsed", [False, True])
def test_cli_extracts_collage_and_keeps_inputs_unchanged(tmp_path: Path, collapsed: bool) -> None:
    book = Workbook()
    book.active.append(["SKU", "Supplier", "Season"])
    book.active.append(["A", "Vendor", "Весна 2027"])
    sheet = book.create_sheet("Images")
    sheet.append(["SKU", "Photo", "Name"])
    sheet.append(["A", None, "First"])
    for color in ("red", "blue"):
        buffer = io.BytesIO()
        Image.new("RGB", (40, 20), color).save(buffer, format="PNG")
        sheet.add_image(ExcelImage(buffer), "B2")
    source = tmp_path / "input.xlsx"
    book.save(source)
    if collapsed:
        # Add a collapsed duplicate alongside the two visible pictures. Its bytes are
        # valid; treating it as a third image would silently alter the collage.
        with ZipFile(source) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        drawing = ET.fromstring(parts["xl/drawings/drawing1.xml"])
        duplicate = ET.fromstring(ET.tostring(drawing[0]))
        duplicate.find(f"{{{D}}}ext").set("cy", "0")
        drawing.append(duplicate)
        parts["xl/drawings/drawing1.xml"] = ET.tostring(drawing)
        with ZipFile(source, "w") as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
    before = source.read_bytes()
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({
        "schema_version": 1, "stage": "mail_capture", "mails": [{
            "entry_id": "entry", "store_id": "store", "subject": "Subject",
            "sender_name": "", "sender_address": "",
            "received_at": datetime(2026, 10, 9, tzinfo=UTC).isoformat(),
        }], "attachments": [{"mail_index": 0, "attachment_index": 1,
            "original_name": source.name, "path": source.name,
            "sha256": sha256(before).hexdigest()}],
    }), encoding="utf-8")
    runtime = tmp_path / "runtime"
    assert main(["preview-photos", "--config", str(CONFIG), "--manifest", str(manifest),
                 "--data-dir", str(runtime)]) == 0
    data = json.loads(next((runtime / "reports").glob("photo-preview-*.json")).read_text('utf-8'))
    assert data["stage"] == "photo_preview"
    assert data["error_count"] == 0
    warnings = [i for i in data["files"][0]["issues"] if i["code"] == "PHOTO_NON_RENDERING"]
    assert len(warnings) == int(collapsed)
    assert all(i["severity"] == "warning" for i in warnings)
    assert data["photo_count"] == 1 and data["eligible_photo_count"] == 1
    product = data["files"][0]["products"][0]
    asset = product["photo"]
    assert (asset["width"], asset["height"]) == (1024, 512)
    assert asset["image_count"] == 2
    assert Path(asset["path"]).is_file()
    assert product["routing"]["route_id"] == "spring"
    assert source.read_bytes() == before


def test_one_cell_anchor_detects_visible_row_span(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["SKU", "Photo", "Name"])
    sheet.append(["A", None, "First"])
    sheet.append(["B", None, "Second"])
    buffer = io.BytesIO()
    Image.new("RGB", (40, 35), "red").save(buffer, format="PNG")
    sheet.add_image(ExcelImage(buffer), "B2")
    path = tmp_path / "spanning.xlsx"
    book.save(path)
    result = PhotoReader().read(path)
    assert result.pictures[0].row == 2 and result.pictures[0].end_row == 3
