from pathlib import Path

import pytest
from openpyxl import Workbook

from kuchenland_importer.domain.master import (
    ExistingProduct,
    ExistingRow,
    MasterSnapshot,
    SheetLayout,
)
from kuchenland_importer.domain.photos import PhotoAction
from kuchenland_importer.infrastructure.excel.master_writer import WrittenProduct
from kuchenland_importer.infrastructure.excel.verify_master import verify_inventory


@pytest.mark.parametrize("damage", [None, "missing", "extra", "duplicate"])
def test_inventory_verifier_keeps_unmodified_articles_and_accepts_season_move(
    tmp_path: Path, damage
):
    names = ("Old", "Old Photos", "New", "New Photos")
    snapshot = MasterSnapshot(
        {name: SheetLayout(name, 1, 1, 2, {"article": 1}, 3) for name in names},
        {
            "KEEP": ExistingProduct([ExistingRow("Old", 2), ExistingRow("Old Photos", 2)]),
            "MOVE": ExistingProduct([ExistingRow("Old", 3), ExistingRow("Old Photos", 3)]),
        },
        names,
    )
    book = Workbook()
    book.remove(book.active)
    for name in names:
        sheet = book.create_sheet(name)
        sheet.append(["Article"])
        sheet.append(["KEEP" if name.startswith("Old") else "MOVE"])
    if damage == "missing":
        book["Old"].cell(2, 1, "OVERWRITTEN")
    elif damage == "extra":
        book["New"].append(["UNEXPECTED"])
    elif damage == "duplicate":
        book["Old"].append(["KEEP"])
    path = tmp_path / "master.xlsx"
    book.save(path)
    book.close()
    written = (WrittenProduct("MOVE", "New", 2, "New Photos", 2, None, PhotoAction.NONE),)
    if damage:
        with pytest.raises(ValueError, match="сохранности артикулов"):
            verify_inventory(path, snapshot, written)
    else:
        verify_inventory(path, snapshot, written)
