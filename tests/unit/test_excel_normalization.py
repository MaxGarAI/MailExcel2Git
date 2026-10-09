import json
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path

import pytest
from openpyxl import Workbook

from kuchenland_importer.application.normalize_workbook import NormalizeWorkbook, article_value
from kuchenland_importer.application.value_normalization import normalize_value
from kuchenland_importer.domain.cell_values import ExcelErrorValue
from kuchenland_importer.domain.mail import MailMetadata, SavedAttachment
from kuchenland_importer.domain.workbook import SourceCell
from kuchenland_importer.infrastructure.capture_loader import load_attachments
from kuchenland_importer.infrastructure.excel.columns import load_columns
from kuchenland_importer.infrastructure.excel.reader import WorkbookReader
from kuchenland_importer.infrastructure.preview_report import encode_value

CONFIG = Path(__file__).resolve().parents[2] / "config" / "columns.toml"


def attached(path: Path) -> SavedAttachment:
    mail = MailMetadata("entry", "store", "Subject", datetime.now(UTC), "", "")
    return SavedAttachment(mail, 1, path.name, path, sha256(path.read_bytes()).hexdigest())


def preview(path: Path):
    return NormalizeWorkbook(load_columns(CONFIG)).execute(
        attached(path), WorkbookReader().read(path)
    )


def test_reordered_synonyms_unknown_fields_and_photo_join(tmp_path: Path) -> None:
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    main = workbook.active
    main.append(["Season", "Supplier", "Unit Price", "SKU", "Custom field"])
    main.append(["Весна 2027", "Vendor", "1 234,50", "001-A", "Keep me"])
    main.append(["Лето 2027", "Vendor", 0, "B", "Other"])
    photo = workbook.create_sheet("Images")
    photo.append(["Артикул", "Фото", "Описание"])
    photo.append(["B", None, "Second"])
    photo.append(["001-A", None, "First"])
    workbook.save(path)
    before = path.read_bytes()
    result = preview(path)
    assert not result.issues
    assert len(result.products) == 2
    first = result.products[0]
    assert first.article == "001-A"
    assert first.values["purchase_price"] == Decimal("1234.50")
    assert first.values["source:custom field"] == "Keep me"
    assert first.values["photo:description"] == "First"
    assert result.products[1].values["purchase_price"] == Decimal("0")
    assert path.read_bytes() == before


def test_multilevel_header_with_merged_groups(tmp_path: Path) -> None:
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Артикул", "Поставщик", "Сезон", "FOB", None, "Наименование"])
    sheet.append([None, None, None, "ПРЕДЫДУЩИЙ", "НОВЫЙ", None])
    for col in ("A", "B", "C", "F"):
        sheet.merge_cells(f"{col}1:{col}2")
    sheet.merge_cells("D1:E1")
    sheet.append(["A", "Vendor", "Зима", 5, 6, "Name"])
    book.save(path)
    result = preview(path)
    assert len(result.products) == 1
    assert result.products[0].values["purchase_price"] == Decimal("6")
    assert result.products[0].source_row == 3


def test_excel_error_remains_typed_and_formula_without_cache_fails(tmp_path: Path) -> None:
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Артикул", "Поставщик", "Сезон", "Цена"])
    sheet.append(["A", "Vendor", "Весна", "#DIV/0!"])
    sheet.append(["B", "Vendor", "Весна", "=1+1"])
    book.save(path)
    result = preview(path)
    assert len(result.products) == 1
    assert result.products[0].values["purchase_price"] == ExcelErrorValue("#DIV/0!")
    assert {issue.code for issue in result.issues} == {"EXCEL_ERROR_VALUE", "ROW_INVALID"}
    assert encode_value(ExcelErrorValue("#DIV/0!")) == {"type": "excel_error", "code": "#DIV/0!"}


def test_cached_empty_string_is_not_missing_formula_cache(tmp_path: Path) -> None:
    path = tmp_path / "empty-result.xlsx"
    book = Workbook()
    book.active.append(["Артикул", "Поставщик", "Сезон", "Optional"])
    book.active.append(["A", "Vendor", "Весна", '=IF(1=1,"","")'])
    book.save(path)
    # Produce the OOXML cached empty-string representation saved by Excel.
    with zipfile.ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    contents["xl/worksheets/sheet1.xml"] = contents["xl/worksheets/sheet1.xml"].replace(
        b'<c r="D2">', b'<c r="D2" t="str">'
    )
    with zipfile.ZipFile(path, "w") as archive:
        for name, value in contents.items():
            archive.writestr(name, value)
    result = preview(path)
    assert len(result.products) == 1
    assert not result.issues
    assert result.products[0].values["source:optional"] == ""


@pytest.mark.parametrize("text", ["SKU", "Article", "Артикул", "Item No."])
def test_article_aliases(text: str) -> None:
    assert load_columns(CONFIG).resolve(text) == "article"


def test_leading_zeros_and_no_fractional_article() -> None:
    assert article_value(SourceCell(12, "00000")) == "00012"
    assert article_value(SourceCell("00012")) == "00012"
    with pytest.raises(ValueError):
        article_value(SourceCell(12.5))


def test_ambiguous_sku_columns_are_rejected(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    book = Workbook()
    book.active.append(["Артикул", "SKU", "Поставщик", "Сезон"])
    book.active.append(["A", "B", "Vendor", "Весна"])
    book.save(path)
    result = preview(path)
    assert not result.products
    assert any(issue.code == "SHEET_SCHEMA" for issue in result.issues)


def test_no_supplier_guessing_from_sender(tmp_path: Path) -> None:
    path = tmp_path / "no-vendor.xlsx"
    book = Workbook()
    book.active.append(["Артикул", "Поставщик", "Сезон"])
    book.active.append(["A", None, "Весна"])
    book.save(path)
    assert not preview(path).products


def test_manifest_path_cannot_escape_capture_directory(tmp_path: Path) -> None:
    path = tmp_path / "manifest.json"
    path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "stage": "mail_capture",
                "mails": [],
                "attachments": [{"path": "../outside.xlsx"}],
            }
        ),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="выходит"):
        load_attachments(path)


def test_date_normalization_without_inventing_year() -> None:
    assert normalize_value("sales_start", "03.01.2027") == date(2027, 1, 3)
    assert normalize_value("sales_start", "декабрь") == "декабрь"
